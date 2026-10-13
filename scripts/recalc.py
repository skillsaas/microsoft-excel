#!/usr/bin/env python3
"""Force a full recalculation of a workbook with headless LibreOffice, then verify.

Usage:
  python recalc.py INPUT.xlsx OUTDIR [--timeout 180] [--soffice PATH] [--force] [--sandboxed]

Writes OUTDIR/<input name> with fresh cached values. Never touches INPUT, and
refuses to replace an existing OUTDIR/<input name> unless --force is given.
A plain LibreOffice load may keep the cached values stored in the file; this
script uses a throwaway profile that sets "recalculate on load: always",
"update links: never" and disables macros, then reports formulas still missing
a cached value and error cells.

Network: a forced recalc evaluates WEBSERVICE and similar functions, so a
hostile workbook could send cell data to the internet. On Linux the script runs
LibreOffice inside a new network namespace with no network (`unshare -n` or
`unshare -rn`). Where that is unavailable (macOS, Windows, locked-down
containers, or Ubuntu 24.04+ where AppArmor blocks unprivileged user
namespaces) it refuses to run.

--sandboxed is NOT a way past that refusal. It is an operator's statement that
this process already runs in a sandbox with no network (for example
`docker run --network none`). An AI agent must never add it on its own
judgment; only the human or operator who set up that sandbox may pass it. If
the script refuses, stop and tell the user. Even with --sandboxed the script
first tries short TCP connects to public addresses and refuses (exit 2) if any
succeeds.

Local files: a network namespace does not hide the filesystem. Link updates are
set to never, the input is copied into a private temp folder and LibreOffice
runs there with HOME and the working directory set to that folder, so relative
paths do not reach the user's folders. In our tests WEBSERVICE("file://..."),
FILTERXML and INDIRECT of a file:// target returned errors, not file contents,
but other paths may exist. For hostile files use a container or bwrap that
mounts only the input (references/security.md).

.xlsm: LibreOffice rewrites the whole package. The VBA project, form controls
and ActiveX parts may not survive byte for byte. The script reports whether
xl/vbaProject.bin is identical, changed or missing. Recalculate a copy, keep
the original, and never hand the recalculated .xlsm back as the macro master.

--force replaces OUTDIR/<input name> without a backup.

Limits: LibreOffice is a different engine from Excel. GROUPBY, PIVOTBY, REGEX*
and some LAMBDA helpers return #NAME?. Say which engine produced the numbers.
A cached value being present does not prove the result matches Excel.
Exit codes: 0 ok, 1 recalculated but errors, missing values or a changed VBA
project, 2 failure or refusal (including a working network under --sandboxed).
Requires: LibreOffice (soffice), pip install openpyxl defusedxml
"""
import argparse
import hashlib
import os
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
from pathlib import Path

XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>1</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="BlockUntrustedRefererLinks" oor:op="fuse"><value>true</value></prop></item>
</oor:items>
"""
# Calc Content/Update/Link: 0 always, 1 never, 2 on request (LibreOffice default 2).
ERRORS = {"#N/A", "#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#NUM!", "#NULL!", "#SPILL!", "#CALC!", "Err:502", "Err:504"}
FILTERS = {".xlsx": "xlsx:Calc MS Excel 2007 XML", ".xlsm": "xlsm:Calc MS Excel 2007 VBA XML"}


def die(msg):
    """Refusals and setup errors exit 2, like other failures."""
    sys.stderr.write(msg + "\n")
    sys.exit(2)


def find_soffice(explicit):
    for c in [explicit, shutil.which("soffice"), shutil.which("libreoffice"),
              "/Applications/LibreOffice.app/Contents/MacOS/soffice",
              r"C:\Program Files\LibreOffice\program\soffice.exe"]:
        if c and Path(c).exists():
            return c
    return None


def no_network_prefix():
    """Return a command prefix that runs a child with no network, or None."""
    unshare = shutil.which("unshare")
    if not unshare or not sys.platform.startswith("linux"):
        return None
    for prefix in ([unshare, "-n"], [unshare, "-rn"]):
        try:
            if subprocess.run(prefix + ["true"], capture_output=True, timeout=10).returncode == 0:
                return prefix
        except (OSError, subprocess.TimeoutExpired):
            pass
    return None


# Public addresses tried under --sandboxed. XLRECALC_EXTRA_PROBES (host:port,...)
# can add targets, for tests; it cannot remove these.
PROBES = [("1.1.1.1", 443), ("8.8.8.8", 53), ("9.9.9.9", 443)]


def network_reachable(timeout=1.0):
    """Return the first host:port that accepts a TCP connect, or None."""
    targets = list(PROBES)
    for item in filter(None, os.environ.get("XLRECALC_EXTRA_PROBES", "").split(",")):
        host, _, port = item.rpartition(":")
        if host and port.isdigit():
            targets.append((host, int(port)))
    for host, port in targets:
        try:
            with socket.create_connection((host, port), timeout=timeout):
                return f"{host}:{port}"
        except OSError:
            pass
    return None


def vba_digest(path):
    """sha256 of xl/vbaProject.bin, or None if the package has none."""
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            return hashlib.sha256(z.read("xl/vbaProject.bin")).hexdigest()
    except (KeyError, zipfile.BadZipFile, OSError):
        return None


def vba_status(before, after):
    if before is None:
        return "none"
    if after is None:
        return "missing"
    return "identical" if before == after else "changed"


def verify(path):
    """Count formulas, formulas without a cached value, and error cells.
    Cells are matched by coordinate, so differing row lengths cannot hide cells."""
    import defusedxml  # noqa: F401  safe XML parsing in openpyxl
    from openpyxl import load_workbook
    wf = load_workbook(path, read_only=True)
    wv = load_workbook(path, read_only=True, data_only=True)
    missing = errors = formulas = 0
    for ws in wf.worksheets:
        values = {}
        for row in wv[ws.title].iter_rows():
            for c in row:
                if getattr(c, "row", None) is not None and c.value is not None:
                    values[(c.row, c.column)] = c.value
        for row in ws.iter_rows():
            for c in row:
                if getattr(c, "row", None) is None:
                    continue
                v = values.get((c.row, c.column))
                if c.data_type == "f":
                    formulas += 1
                    if v is None:
                        missing += 1
                if isinstance(v, str) and v in ERRORS:
                    errors += 1
    wf.close()
    wv.close()
    return formulas, missing, errors


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("outdir")
    ap.add_argument("--timeout", type=int, default=180)
    ap.add_argument("--soffice")
    ap.add_argument("--force", action="store_true",
                    help="replace an existing OUTDIR/<input name> (no backup is kept)")
    ap.add_argument("--sandboxed", action="store_true",
                    help="OPERATOR ONLY: run without unshare because a human or operator already put this "
                         "process in a no-network sandbox. Agents must never add this flag on their own to get "
                         "past a refusal. The script still refuses if a test connection succeeds.")
    a = ap.parse_args()

    src = Path(a.input).resolve()
    out = Path(a.outdir).resolve()
    if not src.is_file():
        die(f"error: {src} not found")
    if src.suffix.lower() not in FILTERS:
        die("error: only .xlsx and .xlsm are supported")
    if out == src.parent:
        die("error: OUTDIR must differ from the input folder so the original is never overwritten")
    dest = out / src.name
    if dest.exists() and not a.force:
        die(f"error: {dest} already exists; choose another OUTDIR or pass --force")
    soffice = find_soffice(a.soffice)
    if not soffice:
        die("error: LibreOffice not found (install it or pass --soffice)")
    if a.sandboxed:
        reached = network_reachable()
        if reached:
            die(f"error: --sandboxed was passed but a test connection to {reached} succeeded, so this "
                     "process has network access. Refusing. Do not retry; tell the user.")
        prefix = []
    else:
        prefix = no_network_prefix()
    if prefix is None:
        die("error: cannot disable network for LibreOffice here (unshare is missing, not Linux, or user "
                 "namespaces are blocked, for example by AppArmor on Ubuntu 24.04+). Do not add --sandboxed "
                 "yourself: stop and tell the user. A human or operator can run this inside a sandbox with no "
                 "network (for example docker run --network none) and pass --sandboxed")
    before_vba = vba_digest(src) if src.suffix.lower() == ".xlsm" else None

    with tempfile.TemporaryDirectory(prefix="xlrecalc-") as prof:
        (Path(prof) / "user").mkdir()
        (Path(prof) / "user" / "registrymodifications.xcu").write_text(XCU, encoding="utf-8")
        tmp_out = Path(prof) / "out"
        tmp_out.mkdir()
        # Work on a private copy so relative link targets resolve inside the temp folder,
        # not next to the user's file.
        work = Path(prof) / "in"
        work.mkdir()
        copy = work / src.name
        shutil.copyfile(src, copy)
        cmd = prefix + [soffice, f"-env:UserInstallation={Path(prof).as_uri()}", "--headless", "--norestore",
                        "--nolockcheck", "--convert-to", FILTERS[src.suffix.lower()],
                        "--outdir", str(tmp_out), str(copy)]
        env = {"PATH": os.environ.get("PATH", ""), "HOME": prof, "TMPDIR": prof,
               "LANG": os.environ.get("LANG", "C.UTF-8"), "SAL_USE_VCLPLUGIN": "svp"}
        if os.name == "nt":
            env.update({k: os.environ[k] for k in ("SYSTEMROOT", "TEMP", "TMP") if k in os.environ})
        try:
            p = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, env=env,
                                 cwd=str(work), start_new_session=(os.name != "nt"))
        except OSError as e:
            sys.stderr.write(f"error: could not start LibreOffice: {e}\n")
            return 2
        try:
            stdout, stderr = p.communicate(timeout=a.timeout)
        except subprocess.TimeoutExpired:
            if os.name != "nt":
                try:
                    os.killpg(p.pid, signal.SIGKILL)  # also kills soffice.bin children
                except OSError:
                    pass
            p.kill()
            p.communicate()
            sys.stderr.write(f"error: LibreOffice timed out after {a.timeout}s and was killed\n")
            return 2
        produced = tmp_out / src.name
        if p.returncode != 0 or not produced.exists():
            sys.stderr.write((stdout or "") + (stderr or ""))
            return 2
        out.mkdir(parents=True, exist_ok=True)
        shutil.move(str(produced), str(dest))  # only now is an existing file replaced (with --force)

    formulas, missing, errors = verify(dest)
    mode = "operator sandbox, connect test failed as expected" if a.sandboxed else "disabled"
    print(f"recalculated with LibreOffice (network {mode}) -> {dest}")
    print(f"formulas={formulas} missing_cached_value={missing} error_cells={errors}")
    vba = "n/a"
    if src.suffix.lower() == ".xlsm":
        vba = vba_status(before_vba, vba_digest(dest))
        print(f"vba_project={vba}")
        print("note: keep the original .xlsm as the macro master; use this copy only to read values")
    return 1 if (missing or errors or vba in ("changed", "missing")) else 0


if __name__ == "__main__":
    sys.exit(main())
